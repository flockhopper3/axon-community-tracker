"""Portable GitHub Actions exports; values are computed from preserved snapshots."""
from __future__ import annotations

import csv
import json
from pathlib import Path

from .model import DIRECTORY_URL, METRICS, PRIMARY, valid_count


def json_file(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf-8")


def text_value(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return value


def csv_safe(value):
    value = text_value(value)
    # Upstream strings are untrusted. Avoid spreadsheet formula injection.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def columns(rows, default=()):
    return list(dict.fromkeys([*default, *(k for row in rows for k in row)]))


def csv_file(path, rows, default=()):
    path.parent.mkdir(parents=True, exist_ok=True)
    keys = columns(rows, default)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=keys)
        writer.writeheader()
        writer.writerows({k: csv_safe(v) for k, v in row.items()} for row in rows)


def md(value):
    return str(value).replace("|", "\\|").replace("\n", " ").replace("<", "&lt;")


def report(manifest, orgs, events, ranked):
    stamp = manifest["captured_at"]
    lines = [f"# Camera statistics — {stamp}", "",
        f"Directory entries: **{manifest['directory_entries']}**. Distinct organizations: **{len(orgs)}**. "
        f"Complete statistics: **{manifest['complete_orgs']} ({manifest['coverage']:.1%})**.", "",
        "Counts are publicly reported API values, not independently verified camera inventories. "
        "Different count categories overlap. Sums across organizations may also include the same physical cameras.", ""]
    if not manifest["quality_ok"]:
        lines += ["## Collection needs review", "", *["- " + reason for reason in manifest["quality_reasons"]], "",
                  "This run is preserved but excluded from future comparison baselines. Growth rankings are withheld.", ""]
    lines += ["## Reported counts", "", "| Metric | Sum across distinct organizations | Organizations reporting |",
              "|---|---:|---:|"]
    for metric in METRICS:
        values = [r[metric] for r in orgs if valid_count(r.get(metric))]
        total = f"{sum(values):,}" if values else "Unavailable"
        lines.append(f"| {metric} | {total} | {len(values)} / {len(orgs)} |")
    lines += ["", "## Largest reported integrated-camera inventories", "", "| Program | State | Integrated |", "|---|---|---:|"]
    for row in sorted((r for r in orgs if valid_count(r.get(PRIMARY))), key=lambda r: r[PRIMARY], reverse=True)[:15]:
        lines.append(f"| {md(row['programs'])} | {md(row['state'])} | {row[PRIMARY]:,} |")
    labels = {"fastest_absolute": "Largest increases", "fastest_weekly_rate": "Fastest increases per week",
              "fastest_percentage": "Largest percentage increases (baseline at least 25)", "largest_declines": "Largest decreases"}
    for kind, title in labels.items():
        lines += ["", "## " + title, ""]
        selected = [r for r in ranked if r["ranking"] == kind and r["period"] == "previous" and r["metric"] == PRIMARY][:15]
        if not selected:
            lines += ["No eligible comparisons yet. Rankings require two valid observations at least six days apart; "
                      "the first run establishes a baseline."]
            continue
        lines += ["| Program | State | Before | Now | Change | Change % | Days | Per week |", "|---|---|---:|---:|---:|---:|---:|---:|"]
        for row in selected:
            percent = f"{row['change_pct']:.1%}" if row["change_pct"] is not None else "New from zero"
            lines.append(f"| {md(row['programs'])} | {md(row['state'])} | {row['previous_count']:,} | {row['current_count']:,} | "
                         f"{row['change']:+,} | {percent} | {row['elapsed_days']:.1f} | {row['change_per_week']:+,.1f} |")
    lines += ["", "## Change log", "", "| Event | Organization | Field | Before | After |", "|---|---|---|---|---|"]
    for event in events[:100]:
        def short(value):
            value = text_value(value)
            return md(value if value is not None else "Missing")[:180]
        lines.append(f"| {md(event['event'])} | {md(event['org'])} | {md(event['field'])} | {short(event['previous'])} | {short(event['current'])} |")
    if len(events) > 100:
        lines += [f"", f"Showing 100 of {len(events)} events; the CSV and JSON contain every event."]
    lines += ["", "## Reading the data", "",
        "- Growth is a change in reported counts, not proof of newly installed cameras.",
        "- Failed requests and invalid/missing fields stay blank. Genuine zero values stay zero.",
        "- Directory additions/removals are tracked separately from camera-count changes.",
        "- Repeated directory entries for the same org are fetched once and counted once in summaries.",
        "- Percentage change from zero is undefined. Percentage rankings require a baseline of at least 25.",
        "- Weekly rates use actual elapsed time. Missing observations can extend the comparison interval.",
        "- 28- and 91-day comparisons use the most recent valid observation on or before the target date.",
        "- State groups use program directory labels, not camera locations or guaranteed US-only geography.",
        "", f"Source directory: {DIRECTORY_URL}", "",
        "See README.md for full methodology and exports. Raw responses and capture metadata are retained with each run.", ""]
    return "\n".join(lines)


def excel_file(path, tables, manifest):
    """openpyxl is installable on hosted runners; the desktop artifact runtime is not."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    overview = wb.active
    overview.title = "Overview"
    overview.append(["Axon Community Connect camera statistics"])
    overview.append(["Captured at (UTC)", manifest["captured_at"]])
    overview.append(["Directory entries", manifest["directory_entries"]])
    overview.append(["Distinct organizations", manifest["distinct_orgs"]])
    overview.append(["Complete statistics coverage", manifest["coverage"]])
    overview.append(["Source", DIRECTORY_URL])
    overview.append(["Interpretation", "Reported counts; categories overlap and agency sums may double-count physical cameras."])
    overview.append(["Growth", "Two observations at least six days apart are needed for rankings. Missing data is not zero."])
    overview.append(["Quality", "Accepted baseline" if manifest["quality_ok"] else "; ".join(manifest["quality_reasons"])])
    overview.append(["Analysis", "Exports are immutable computed snapshots. Re-run the collector to refresh all calculations."])
    overview.append(["Metric", "Reported sum", "Organizations reporting"])
    for metric in METRICS:
        values = [r[metric] for r in tables["Organizations"] if valid_count(r.get(metric))]
        overview.append([metric, sum(values) if values else None, len(values)])
    overview["B5"].number_format = "0.0%"
    for row in range(12, 18):
        overview.cell(row, 2).number_format = "#,##0"
    overview.column_dimensions["A"].width = 34
    overview.column_dimensions["B"].width = 82
    overview.column_dimensions["C"].width = 25
    overview.merge_cells("A1:C1")
    for row in (7, 8, 9, 10):
        overview.row_dimensions[row].height = 34
        overview.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="center")
    for i, (name, rows) in enumerate(tables.items(), 1):
        ws = wb.create_sheet(name)
        keys = columns(rows)
        if not keys:
            ws.append(["No observations available yet"])
            ws.column_dimensions["A"].width = 42
        else:
            ws.append(keys)
            for data in rows:
                ws.append([text_value(data.get(k)) for k in keys])
            # Force strings to text, including strings beginning with '='.
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                    elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        header = keys[cell.column - 1]
                        cell.number_format = "0.000000" if header.endswith((".lat", ".lng")) else "0.0%" if header in ("change_pct", "coverage") else "#,##0.0" if header in ("elapsed_days", "change_per_week") else "#,##0"
            if rows:
                table = Table(displayName=f"DataTable{i}", ref=f"A1:{get_column_letter(len(keys))}{len(rows)+1}")
                table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
                ws.add_table(table)
            ws.freeze_panes = "C2"
            for j, key in enumerate(keys, 1):
                observed_width = max((len(str(row.get(key) or "")) for row in rows), default=0)
                width = max(18, min(52, max(len(key), observed_width) + 3))
                if key.endswith("_at") or key.endswith(".observed_at"):
                    width = 28
                if key in ("programs", "title", "note", "error"):
                    width = min(64, max(36, observed_width + 3))
                ws.column_dimensions[get_column_letter(j)].width = width
                if key in ("programs", "title", "note", "error"):
                    for cell in ws[get_column_letter(j)][1:]:
                        cell.alignment = Alignment(wrap_text=True, vertical="center")
                        if cell.value and len(str(cell.value)) > width:
                            ws.row_dimensions[cell.row].height = max(ws.row_dimensions[cell.row].height or 15, min(90, 15 * (len(str(cell.value)) // int(width) + 1)))
            ws.row_dimensions[1].height = 42
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="183E56")
            cell.font = Font(name="Arial", color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.sheet_view.showGridLines = False
    # One category only: never stack overlapping camera-count categories.
    chart_ws = wb.create_sheet("Largest inventories")
    chart_ws.append(["Program", "Integrated cameras"])
    top = sorted((r for r in tables["Organizations"] if valid_count(r.get(PRIMARY))), key=lambda r: r[PRIMARY], reverse=True)[:15]
    for row in top:
        chart_ws.append([row["programs"], row[PRIMARY]])
        chart_ws.cell(chart_ws.max_row, 2).number_format = "#,##0"
    chart_ws.column_dimensions["A"].width = 52
    chart_ws.column_dimensions["B"].width = 22
    if top:
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Largest reported integrated-camera inventories"
        chart.add_data(Reference(chart_ws, min_col=2, min_row=1, max_row=len(top)+1), titles_from_data=True)
        chart.set_categories(Reference(chart_ws, min_col=1, min_row=2, max_row=len(top)+1))
        chart.width, chart.height = 28, 17
        chart.legend = None
        chart.x_axis.scaling.min = 0
        chart.x_axis.numFmt = '#,##0'
        chart_ws.add_chart(chart, "D2")
    for ws in wb:
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if cell.row != 1:
                    cell.font = Font(name="Arial", size=11)
        ws.sheet_view.showGridLines = False
    for row in (1, 11):
        for cell in overview[row]:
            cell.fill = PatternFill("solid", fgColor="183E56")
            cell.font = Font(name="Arial", color="FFFFFF", bold=True, size=12)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def exports(folder, run_id, manifest, orgs, communities, events, deltas, state_rows, ranked):
    datasets = {"organizations": (orgs, ("snapshot_at", "org", "programs", "city", "state", *METRICS)),
                "communities": (communities, ("snapshot_at", "id", "title", "org")),
                "changes": (events, ("event", "org", "field", "previous", "current", "note")),
                "growth": (deltas, ("org", "metric", "period", "previous_count", "current_count", "change", "change_pct", "change_per_week")),
                "state_summary": (state_rows, ("state", "metric", "reported_sum", "orgs_with_metric")),
                "rankings": (ranked, ("ranking", "rank", "org", "metric", "period", "change"))}
    for name, (rows, default) in datasets.items():
        csv_file(folder / f"{name}_{run_id}.csv", rows, default)
        json_file(folder / f"{name}_{run_id}.json", rows)
    excel_file(folder / f"camera_statistics_{run_id}.xlsx",
               {"Organizations": orgs, "Communities": communities, "Growth": deltas,
                "Rankings": ranked, "State summary": state_rows, "Changes": events}, manifest)
    (folder / f"report_{run_id}.md").write_text(report(manifest, orgs, events, ranked), encoding="utf-8")
