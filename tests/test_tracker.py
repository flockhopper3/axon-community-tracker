import argparse
import csv
import datetime as dt
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import urllib.error

from axon_tracker.collect import PublicClient
from axon_tracker.export import csv_file, excel_file
from axon_tracker.model import (METRICS, PRIMARY, changes, growth, make_communities,
                                make_orgs, quality, rankings, summaries, validate_directory)
from axon_tracker.__main__ import run


def listing(org="boulderpdco", id="boulder", state="CO"):
    return {"id": id, "org": org, "title": id, "url": "https://example.org", "location": {"city": id, "state": state}}


def response(count=100, timestamp="2026-01-01T00:00:00Z"):
    return {"payload": {metric: count for metric in METRICS}, "retrieved_at": timestamp, "http_status": 200}


def snapshot(count=100, timestamp="2026-01-01T00:00:00Z", locations=None, quality_ok=True):
    locations = locations or [listing()]
    rows = make_orgs(locations, {x["org"]: response(count, timestamp) for x in locations}, timestamp)
    return {"schema_version": 1, "manifest": {"quality_ok": quality_ok, "captured_at": timestamp, "run_id": timestamp},
            "locations": locations, "orgs": rows}


class AnalysisTests(unittest.TestCase):
    def test_duplicate_org_preserved_in_directory_but_counted_once(self):
        locations = [listing(), listing(id="second")]
        rows = make_orgs(locations, {"boulderpdco": response(161)}, "2026-01-01T00:00:00Z")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["program_count"], 2)
        self.assertEqual(len(make_communities(locations, rows, "now")), 2)
        self.assertEqual(summaries(rows, [])[1]["reported_sum"], 161)

    def test_failed_response_is_not_zero(self):
        rows = make_orgs([listing()], {"boulderpdco": {"error": "HTTP 503", "http_status": 503}}, "2026-01-08T00:00:00Z")
        self.assertIsNone(rows[0][PRIMARY])
        self.assertEqual(rows[0]["status"], "error")
        self.assertEqual(growth(rows, [snapshot()]), [])

    def test_invalid_counts_and_schema(self):
        r = response()
        r["payload"].update(totalIntegratedCameras=-1, subscribedCameras=True, totalOwnedCameras="100")
        rows = make_orgs([listing()], {"boulderpdco": r}, "2026-01-01T00:00:00Z")
        self.assertEqual(rows[0]["status"], "partial")
        self.assertIsNone(rows[0][PRIMARY])
        self.assertIsNone(rows[0]["subscribedCameras"])
        self.assertIsNone(rows[0]["totalOwnedCameras"])

    def test_true_zero_stays_zero(self):
        rows = snapshot(0)["orgs"]
        self.assertEqual(rows[0][PRIMARY], 0)
        self.assertEqual(rows[0]["status"], "ok")

    def test_actual_elapsed_time_and_percentage(self):
        previous = snapshot(100)
        rows = snapshot(140, "2026-01-15T00:00:00Z")["orgs"]
        delta = next(r for r in growth(rows, [previous]) if r["metric"] == PRIMARY)
        self.assertEqual(delta["change"], 40)
        self.assertEqual(delta["change_per_week"], 20)
        self.assertAlmostEqual(delta["change_pct"], .4)

    def test_zero_baseline_has_no_percentage(self):
        rows = growth(snapshot(10, "2026-01-08T00:00:00Z")["orgs"], [snapshot(0)])
        self.assertIsNone(rows[0]["change_pct"])
        self.assertTrue(rows[0]["from_zero"])
        self.assertFalse(any(r["ranking"] == "fastest_percentage" for r in rankings(rows)))

    def test_short_interval_excluded_from_rankings(self):
        deltas = growth(snapshot(200, "2026-01-01T01:00:00Z")["orgs"], [snapshot()])
        self.assertTrue(deltas)
        self.assertEqual(rankings(deltas), [])

    def test_last_valid_observation_after_missing_week(self):
        bad = snapshot(999, "2026-01-08T00:00:00Z", quality_ok=False)
        deltas = growth(snapshot(120, "2026-01-15T00:00:00Z")["orgs"], [snapshot(), bad])
        self.assertEqual(deltas[0]["previous_count"], 100)
        self.assertEqual(deltas[0]["elapsed_days"], 14)

    def test_28_day_baseline_on_or_before_target(self):
        history = [snapshot(10), snapshot(15, "2026-01-08T00:00:00Z")]
        deltas = growth(snapshot(30, "2026-01-30T00:00:00Z")["orgs"], history)
        row = next(x for x in deltas if x["metric"] == PRIMARY and x["period"] == "28d")
        self.assertEqual(row["previous_count"], 10)
        self.assertEqual(row["elapsed_days"], 29)

    def test_first_observation_not_growth(self):
        self.assertEqual(growth(snapshot()["orgs"], []), [])
        self.assertEqual(changes([listing()], snapshot()["orgs"], None, [])[0]["event"], "baseline")

    def test_directory_removal_not_camera_zero(self):
        before = snapshot(locations=[listing(), listing("other", "other")])
        events = changes([listing()], snapshot()["orgs"], before, [before])
        self.assertEqual([r["event"] for r in events], ["listing_removed"])

    def test_reappearance_and_org_remapping(self):
        first = snapshot(locations=[listing(), listing("other", "other")])
        previous = snapshot()
        events = changes(first["locations"], first["orgs"], previous, [first, previous])
        self.assertIn("listing_reappeared", [r["event"] for r in events])
        events = changes([listing("neworg")], snapshot(locations=[listing("neworg")])["orgs"], previous, [previous])
        self.assertEqual(events[0]["event"], "listing_changed")

    def test_preserve_unknown_nested_data(self):
        locations = [listing() | {"new_field": {"array": [1, 2]}}]
        r = response()
        r["payload"]["futureMetric"] = 42
        rows = make_orgs(locations, {"boulderpdco": r}, "2026-01-01T00:00:00Z")
        communities = make_communities(locations, rows, "now")
        self.assertEqual(communities[0]["new_field.array"], "[1, 2]")
        self.assertEqual(communities[0]["stats.futureMetric"], 42)

    def test_missing_or_malicious_org_not_fetched(self):
        locations = [listing(), listing("../other", "bad"), listing(None, "missing")]
        rows = make_orgs(locations, {"boulderpdco": response()}, "now")
        self.assertEqual(len(rows), 1)
        communities = make_communities(locations, rows, "now")
        self.assertEqual(communities[1]["collection.status"], "invalid_org")
        self.assertEqual(communities[2]["collection.status"], "missing_org")

    def test_quality_and_empty_directory(self):
        with self.assertRaises(ValueError):
            validate_directory([])
        previous = snapshot(locations=[listing(str(i), str(i)) for i in range(10)])
        current = snapshot()
        self.assertFalse(quality(current["locations"], current["orgs"], previous)["quality_ok"])

    def test_percent_ranking_excludes_small_baselines(self):
        deltas = growth(snapshot(500, "2026-01-08T00:00:00Z")["orgs"], [snapshot(5)])
        self.assertTrue(rankings(deltas))
        self.assertFalse(any(r["ranking"] == "fastest_percentage" for r in rankings(deltas)))


class ExportTests(unittest.TestCase):
    def test_csv_formula_injection_and_missing_values(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "test.csv"
            csv_file(path, [{"name": "=HYPERLINK(\"bad\")", "missing": None, "zero": 0}])
            with path.open(encoding="utf-8-sig") as stream:
                row = next(csv.DictReader(stream))
            self.assertTrue(row["name"].startswith("'="))
            self.assertEqual(row["missing"], "")
            self.assertEqual(row["zero"], "0")

    def test_excel_typed_counts_text_and_chart(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "test.xlsx"
            snap = snapshot()
            snap["orgs"][0]["programs"] = "=1+1"
            manifest = {"captured_at": "now", "directory_entries": 1, "distinct_orgs": 1, "coverage": 1., "quality_ok": True}
            excel_file(path, {"Organizations": snap["orgs"], "Growth": []}, manifest)
            wb = load_workbook(path)
            ws = wb["Organizations"]
            headers = {c.value: c.column for c in ws[1]}
            self.assertEqual(ws.cell(2, headers[PRIMARY]).value, 100)
            self.assertEqual(ws.cell(2, headers["programs"]).data_type, "s")
            self.assertEqual(wb["Largest inventories"]["A2"].data_type, "s")
            self.assertEqual(ws.freeze_panes, "C2")
            self.assertEqual(len(wb["Largest inventories"]._charts), 1)


class ClientTests(unittest.TestCase):
    @patch("axon_tracker.collect.time.sleep")
    @patch("axon_tracker.collect.urllib.request.urlopen")
    def test_retry_clears_stale_error(self, urlopen, sleep):
        success = unittest.mock.MagicMock()
        success.__enter__.return_value = success
        success.read.return_value = b'{"count": 1}'
        success.status, success.url, success.headers = 200, "https://example.org", {}
        failure = urllib.error.HTTPError("https://example.org", 503, "unavailable", {}, None)
        urlopen.side_effect = [failure, success]
        with tempfile.TemporaryDirectory() as tmp:
            result = PublicClient(delay=0, retries=1).get("https://example.org", Path(tmp) / "raw.json")
        self.assertNotIn("error", result)
        self.assertEqual(result["attempts"], 2)

    @patch("axon_tracker.collect.urllib.request.urlopen")
    def test_access_denial_is_not_retried(self, urlopen):
        urlopen.side_effect = urllib.error.HTTPError("https://example.org", 403, "denied", {}, None)
        with tempfile.TemporaryDirectory() as tmp:
            result = PublicClient(delay=0).get("https://example.org", Path(tmp) / "raw.json")
        self.assertEqual(urlopen.call_count, 1)
        self.assertEqual(result["http_status"], 403)


class PipelineTests(unittest.TestCase):
    @patch("axon_tracker.__main__.PublicClient")
    def test_end_to_end_and_integrity(self, client):
        def fetch(url, path):
            payload = [listing(), listing(id="second")] if url.endswith("locations.json") else response()["payload"]
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(json.dumps(payload))
            return {"payload": payload, "retrieved_at": "2026-01-01T00:00:00Z", "http_status": 200}
        client.return_value.get.side_effect = fetch
        with tempfile.TemporaryDirectory() as tmp:
            code = run(argparse.Namespace(data_dir=tmp, delay=1, timeout=5, retries=0, min_coverage=.8))
            self.assertEqual(code, 0)
            self.assertEqual(client.return_value.get.call_count, 2)
            manifest = json.loads((Path(tmp)/"latest/manifest.json").read_text())
            root = Path(tmp)/"runs"/manifest["run_id"]
            self.assertTrue((root/"SHA256SUMS.txt").exists())
            self.assertTrue((Path(tmp)/"latest/camera_statistics.xlsx").exists())
            self.assertTrue((Path(tmp)/"history/organizations.csv").exists())


if __name__ == "__main__":
    unittest.main()
